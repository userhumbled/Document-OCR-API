import os
import json
import tempfile
from docx import Document

def export_doc_file(page_results, output_format):
    temp_dir = tempfile.gettempdir()
    
    if output_format == "json":
        json_data = []
        for p_num, content in page_results:
            json_data.append({
                "page_number": p_num,
                "content": content
            })
        display_text = json.dumps(json_data, indent=2, ensure_ascii=False)
        out_path = os.path.join(temp_dir, "extracted_output.json")
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(display_text)
            
    elif output_format == "docx":
        doc = Document()
        display_blocks = []
        for idx, (p_num, content) in enumerate(page_results):
            doc.add_heading(f"[Page {p_num}]", level=1)
            display_blocks.append(f"[Page {p_num}]\n{content}")
            if content.strip():
                for line in content.split("\n"):
                    if line.strip():
                        doc.add_paragraph(line.strip())
            else:
                doc.add_paragraph("[No text detected]")
            if idx < len(page_results) - 1:
                doc.add_page_break()
        out_path = os.path.join(temp_dir, "extracted_output.docx")
        doc.save(out_path)
        display_text = "\n\n".join(display_blocks)
        
    elif output_format == "md":
        sections = []
        for p_num, content in page_results:
            sections.append(f"[Page {p_num}]\n\n{content}")
        display_text = "\n\n---\n\n".join(sections)
        out_path = os.path.join(temp_dir, "extracted_output.md")
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(display_text)
            
    else:
        sections = []
        for p_num, content in page_results:
            sections.append(f"[Page {p_num}]\n{content}")
        display_text = "\n\n".join(sections)
        out_path = os.path.join(temp_dir, "extracted_output.txt")
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(display_text)
            
    return display_text, out_path

def update_api_helper(page_selection, engine_choice, output_format, enable_layout=True, server_url="https://your-public-url.trycloudflare.com"):
    return (
        "from gradio_client import Client, handle_file\n\n"
        f"client = Client('{server_url}')\n"
        "result = client.predict(\n"
        "    handle_file('ocr_demo_test.pdf'),\n"
        f"    '{page_selection}',\n"
        f"    '{engine_choice}',\n"
        f"    '{output_format}',\n"
        f"    {enable_layout},\n"
        "    True,\n"
        "    api_name='/predict'\n"
        ")\n\n"
        "extracted_text = result[0]\n"
        "download_file = result[1]\n"
        "print(extracted_text)\n"
    )
